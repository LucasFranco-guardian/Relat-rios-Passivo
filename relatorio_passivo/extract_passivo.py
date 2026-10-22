#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_passivo.py — lê a planilha FORMATAR_COMPRAxVENDA_PASSIVO_<FUNDO>_v*_MACRO.xlsm (último valor SALVO)
e gera o JSON do Relatório Passivo, usando as mesmas regras de passivo_json.py.

Uso: python3 extract_passivo.py <planilha.xlsm> <saida.json> [--graficos graficos.json]

A planilha NÃO contém os três gráficos por faixa (PARTICIPAÇÃO PF/PJ, CONCENTRAÇÃO e COTISTAS POR FAIXA):
eles nascem no notebook (graf1, qtde_cotas, qtde_cotistas). Passe-os em --graficos (mesmo formato do JSON:
chaves participacao_pf_pj e cotas_faixa.pf_faixa/pj_faixa/cotistas) ou use o caminho 100 % notebook
(passivo_json.montar_json). Sem --graficos os gráficos saem zerados e o script avisa.

Leitura por cabeçalho (sobrevive a linhas inseridas); nunca grava a .xlsm.
"""
import sys, json, argparse, warnings
warnings.filterwarnings("ignore")
import openpyxl, pandas as pd
from passivo_json import (market_maker, top10, cotas_faixa, FAIXAS, _f, _faixa, N_TOP10, N_TOP15)

COLS_DIA = ["corretora", "compra_neg", "compra_vtt", "compra_pct_vtt", "compra_qtt", "compra_med",
            "venda_neg", "venda_vtt", "venda_pct_vtt", "venda_qtt", "venda_med", "saldo_vtt"]
MES = {"GARE11": {"compra": ("MENSAL COMPRA - MKM", 4), "venda": ("HISTÓRICO VENDA - MKM", 4)},   # (aba, coluna pct)
       "GAME11": {"compra": ("HISTÓRICO COMPRA - MKM", 3), "venda": ("HISTÓRICO VENDA - MKM", 3)}}


def achar(ws, texto, max_row=120, max_col=40):
    t = texto.strip().upper()
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for c in row:
            if isinstance(c.value, str) and c.value.strip().upper() == t:
                return c
    raise RuntimeError(f"[{ws.title}] cabeçalho '{texto}' não encontrado")

def bloco_por_cabecalho(ws, ancora, colunas, max_linhas=60):
    """Tabela cujo cabeçalho começa na célula com texto `ancora`; lê até a primeira linha vazia."""
    c = achar(ws, ancora)
    hdr = [ws.cell(c.row, c.column + i).value for i in range(colunas)]
    out = []
    for r in range(c.row + 1, c.row + 1 + max_linhas):
        v = ws.cell(r, c.column).value
        if v is None or str(v).strip() == "": break
        out.append([ws.cell(r, c.column + i).value for i in range(colunas)])
    return pd.DataFrame(out, columns=[str(h).strip() if h is not None else f"c{i}" for i, h in enumerate(hdr)])

def extrair(caminho, graficos=None):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    bt, db = wb["BASE_TABELAS"], wb["DADOS_BROADCAST_COTISTAS"]
    ticker = "GARE11" if "GARE" in caminho.upper() else "GAME11"
    dia = bt.cell(achar(bt, "DATA").row, achar(bt, "DATA").column + 1).value
    dia = pd.to_datetime(dia).date()

    # --- highlights (linha abaixo dos cabeçalhos Preço_D-1 … PJ)
    def abaixo(ws, txt): c = achar(ws, txt); return _f(ws.cell(c.row + 1, c.column).value)
    seg_hdr = "VAR_IFIX_TIJOLO" if ticker == "GARE11" else "VAR_IFIX_PAPEL"
    fech, fech_d1 = abaixo(bt, "Preço_D-1"), abaixo(bt, "Preço_D-2")
    # PREÇO MÉDIO: bloco 'DATA RELATÓRIO | VAR IFIX | PREÇO MÉDIO' da DADOS_BROADCAST (é o que a TABELA_HIGHLIGHT usa)
    c_pm = None
    for row in db.iter_rows(min_row=1, max_row=20, max_col=40):
        for c in row:
            if isinstance(c.value, str) and c.value.strip().upper() == "PREÇO MÉDIO" and c.column > 15: c_pm = c
    if c_pm is None: raise RuntimeError("[DADOS_BROADCAST_COTISTAS] bloco 'DATA RELATÓRIO | VAR IFIX | PREÇO MÉDIO' não encontrado")
    med = _f(db.cell(c_pm.row + 1, c_pm.column).value)
    vol_lbl = "VOLUME TOTAL" if ticker == "GARE11" else "VOLUME FINANCEIRO"
    c_vol = achar(bt, vol_lbl); vol_total = _f(bt.cell(c_vol.row, c_vol.column + 1).value)
    c_vi = achar(db, "VAR IFIX")
    h = {"preco_fech": fech, "preco_max": abaixo(bt, "Maximo"), "preco_min": abaixo(bt, "Minimo"),
         "fech_d1": fech_d1, "preco_med": med,
         "var_pmxpf": (fech - med) / med if med else 0.0, "variacao": (fech - fech_d1) / fech_d1 if fech_d1 else 0.0,
         "var_ifix_seg": abaixo(bt, seg_hdr), "var_ifix": _f(db.cell(c_vi.row + 1, c_vi.column).value),
         "vol_total_mm": vol_total / 1e6, "volume_b3": abaixo(bt, "B3"), "daytrade": abaixo(bt, "DT")}
    if ticker == "GARE11":
        h["taxa_short"] = abaixo(bt, "Taxa Short") / 100.0
        h["vol_short_mm"] = abaixo(bt, "Volume Short") / 1e6

    # --- corretoras do dia + mês anterior
    cd = bloco_por_cabecalho(bt, "CORRETORA", 12); cd.columns = COLS_DIA
    cd["corretora"] = cd["corretora"].astype(str)
    for c in COLS_DIA[1:]: cd[c] = pd.to_numeric(cd[c], errors="coerce").fillna(0.0)
    cm = {}
    for lado, (aba, col_pct) in MES[ticker].items():
        ws = wb[aba]; rows = []
        for r in range(3, 62):
            n = ws.cell(r, 1).value
            if n: rows.append((str(n), _f(ws.cell(r, col_pct).value)))
        cm[lado] = pd.DataFrame(rows, columns=["corretora", "pct"])

    # --- 10+ compradores / vendedores (blocos 'MAIORES COMPRADORES' / 'MAIORES VENDEDORES' → cabeçalho NOME)
    def top_bloco(rotulo):
        a = achar(bt, rotulo)
        for r in range(a.row + 1, a.row + 4):
            if str(bt.cell(r, a.column).value).strip().upper() == "NOME":
                ncol = 8 if ticker == "GARE11" else 6
                hdr = [str(bt.cell(r, a.column + i).value).strip() for i in range(ncol)]
                data = [[bt.cell(rr, a.column + i).value for i in range(ncol)] for rr in range(r + 1, r + 1 + N_TOP10)]
                df = pd.DataFrame(data, columns=hdr).dropna(subset=["NOME"])
                df["VAR_PRECO"] = pd.to_numeric(df["VAR_PRECO"], errors="coerce")
                return df.set_index("NOME")
        raise RuntimeError(f"cabeçalho NOME não achado abaixo de '{rotulo}'")
    tc, tv = top_bloco("MAIORES COMPRADORES"), top_bloco("MAIORES VENDEDORES")

    # --- 15 maiores (cabeçalho NOME | PF_PJ | CATEGORIAS | COTAS)
    t15 = None
    for row in bt.iter_rows(min_row=30, max_row=80, max_col=6):
        for c in row:
            if isinstance(c.value, str) and c.value.strip().upper() == "COTAS":
                r0 = c.row
                data = [[bt.cell(rr, c.column - 3 + i).value for i in range(4)] for rr in range(r0 + 1, r0 + 1 + N_TOP15)]
                t15 = pd.DataFrame(data, columns=["NOME", "PF_PJ", "CATEGORIAS", "COTAS"]).dropna(subset=["NOME"]).set_index("NOME")
    if t15 is None: raise RuntimeError("bloco 15 MAIORES (cabeçalho COTAS) não encontrado")
    b3 = pd.DataFrame({"B3": [abaixo(bt, "B3")], "DT": [abaixo(bt, "DT")], "PF": [abaixo(bt, "PF")], "PJ": [abaixo(bt, "PJ")]})

    # --- cotistas: dia (primeira linha sob 'Entraram') e mês ('TOTAL - CONTINUO')
    c_e, c_s, c_t = achar(db, "Entraram"), achar(db, "Sairam"), achar(db, "TOTAL")
    ent_dia, sai_dia, tot = _f(db.cell(c_e.row + 1, c_e.column).value), _f(db.cell(c_s.row + 1, c_s.column).value), _f(db.cell(c_t.row + 1, c_t.column).value)
    c_cont = achar(db, "TOTAL - CONTINUO")
    mov_mes = {"entrada_mes": _f(db.cell(c_cont.row, c_e.column).value), "saida_mes": _f(db.cell(c_cont.row, c_s.column).value)}

    # --- gráficos (só via --graficos)
    g = graficos or {}
    part = g.get("participacao_pf_pj") or {"faixas": FAIXAS, "venda_pf": [0] * 10, "venda_pj": [0] * 10, "compra_pf": [0] * 10, "compra_pj": [0] * 10}
    cf_g = (g.get("cotas_faixa") or {})
    qc = pd.DataFrame([{"CATEGORIAS": str(i + 1), "PF_PJ": "PF", "% TOTAL": cf_g.get("pf_faixa", [0] * 10)[i]} for i in range(10)] +
                      [{"CATEGORIAS": str(i + 1), "PF_PJ": "PJ", "% TOTAL": cf_g.get("pj_faixa", [0] * 10)[i]} for i in range(10)])
    qt = pd.DataFrame([{"CATEGORIAS": str(i + 1), "% COTISTAS": cf_g.get("cotistas", [0] * 10)[i]} for i in range(10)])
    if not graficos: print("AVISO: gráficos por faixa zerados — passe --graficos ou use passivo_json.montar_json no notebook.", file=sys.stderr)

    return {"fundo": {"ticker": ticker}, "data": dia.strftime("%d/%m/%Y"), "highlights": h,
            "market_maker": market_maker(cd, cm), "participacao_pf_pj": part,
            "vendedores": top10(tv, ticker), "compradores": top10(tc, ticker),
            "cotas_faixa": cotas_faixa(qc, qt, b3, t15, fech, ent_dia, sai_dia, tot, mov_mes)}


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("xlsm"); ap.add_argument("json"); ap.add_argument("--graficos")
    a = ap.parse_args()
    graf = json.load(open(a.graficos, encoding="utf-8")) if a.graficos else None
    d = extrair(a.xlsm, graf)
    json.dump(d, open(a.json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("ok", a.json, d["fundo"]["ticker"], d["data"])

if __name__ == "__main__":
    main()
